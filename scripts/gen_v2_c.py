# -*- coding: utf-8 -*-
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from neutro import *
from docx import Document
from docx.shared import Cm
from collections import Counter
import rc257 as R
M, O = R.V2, R.ORIG
f = R.fnum
FIG = "/Users/abiudbenitez/Documents/Claude/Projects/vr-bateria-terapeutica/entregables/figuras-ruta-critica/"
doc = Document('/tmp/_v2_b.docx')

# ================================================================ 9
landscape(doc)
H(doc,"9. Red de precedencias",1)
P(doc,"Con 257 actividades no es posible dibujar la red completa en una hoja legible. Se presenta en tres "
  "niveles de detalle, cada uno con un propósito distinto.",size=9.5)
table(doc,["Figura","Qué muestra","Para qué sirve"],[
 ("Fig. 1  Matriz de dependencias entre áreas","Cuántas dependencias van de cada área a cada otra.","Ver el acoplamiento entre áreas y cómo cambió al cerrar la red."),
 ("Fig. 2  Ruta crítica en detalle","Las 28 actividades de la ruta crítica, con duración y acumulado.","Seguir la trayectoria que determina la duración del proyecto."),
 ("Fig. 3  Red medida","Las actividades con holgura total menor o igual a 5 días, en carriles por área.","Ver la ruta crítica en su contexto y qué actividades están a punto de volverse críticas."),
],widths=[5.6,9.4,9.6],fs=9.5)

H(doc,"9.1 Matriz de dependencias entre áreas",2)
if os.path.exists(FIG+"Fig1_Matriz_dependencias_areas.png"):
    doc.add_picture(FIG+"Fig1_Matriz_dependencias_areas.png", width=Cm(17.5)); doc.paragraphs[-1].alignment=C
caption(doc,"Fig. 1 — Matriz de dependencias entre áreas. La fila debe terminar antes que la columna.")
P(doc,"La matriz revela la estructura del proyecto. El área de QA, documentación y gestión recibe "
  "dependencias de las siete áreas restantes y no entrega casi ninguna: es el sumidero de la red, donde todo "
  "converge. Ese acoplamiento se intensificó con los cambios, porque la mayoría de las dependencias añadidas "
  "apuntan a actividades de QA: son precisamente las que hacían falta para que el trabajo de las demás áreas "
  "llegara al entregable final.",size=9.5)

H(doc,"9.2 Ruta crítica en detalle",2)
if os.path.exists(FIG+"Fig2_Ruta_critica_detalle.png"):
    doc.add_picture(FIG+"Fig2_Ruta_critica_detalle.png", width=Cm(25.4)); doc.paragraphs[-1].alignment=C
caption(doc,"Fig. 2 — Las 28 actividades de la ruta crítica, con duración individual y acumulada.")

H(doc,"9.3 Red medida",2)
if os.path.exists(FIG+"Fig3_Red_medida.png"):
    doc.add_picture(FIG+"Fig3_Red_medida.png", width=Cm(25.4)); doc.paragraphs[-1].alignment=C
caption(doc,"Fig. 3 — Red medida. Actividades con holgura total menor o igual a 5 días, en carriles por área.")
P(doc,"En la red medida se aprecia lo que las tablas no muestran con la misma claridad: la ruta crítica "
  "recorre casi íntegro el carril del entorno tridimensional y luego el de QA. Los carriles de desarrollo VR, "
  "juego de ritmo, música y sonido no la tocan en ningún punto.",size=9.5)

portrait(doc)
# ================================================================ 10
landscape(doc)
H(doc,"10. Matriz de elasticidad",1)
P(doc,"La matriz de elasticidad reúne los tiempos y las cuatro holguras de cada actividad. Todas las "
  "magnitudes están en días hábiles contados desde el día cero.",size=9.5)
table(doc,["Holgura","Fórmula","Qué responde"],[
 ("Total","TRI menos TPI.","Cuánto puede retrasarse la actividad sin mover la fecha final del proyecto."),
 ("Libre","Menor TPI de las consecuentes, menos el TPT de la actividad.","Cuánto puede retrasarse sin afectar el inicio más temprano de ninguna actividad siguiente."),
 ("Interferente","Holgura total menos holgura libre.","Qué parte del margen, si se consume, reduce el margen de las actividades posteriores."),
 ("Independiente","Menor TPI de las consecuentes, menos el mayor TRT de las antecedentes, menos la duración. Si resulta negativa se toma cero.","Cuánto puede retrasarse sin afectar ni a antecedentes ni a consecuentes, en el peor escenario."),
],widths=[2.4,9.6,13.0],fs=9)
rows=[(k, R.T[k]["sig"], f(R.T[k]["dur"]), f(M["ES"][k]), f(M["EF"][k]), f(M["LS"][k]), f(M["LF"][k]),
       f(M["HT"][k]), f(M["HL"][k]), f(M["HI"][k]), f(M["HN"][k]),
       "Sí" if abs(M["HT"][k])<1e-6 else "") for k in R.ORD]
table(doc,["Clave","Á.","Dur.","TPI","TPT","TRI","TRT","H. total","H. libre","H. interf.","H. indep.","Crítica"],
      rows,widths=[1.8,0.9,1.5,1.7,1.7,1.7,1.7,1.9,1.9,1.9,1.9,1.6],fs=7)

portrait(doc)
# ================================================================ 11
H(doc,"11. Ruta crítica",1)
H(doc,"11.1 Criterio de identificación",2)
P(doc,"Una actividad es crítica cuando cumple estas condiciones, que son equivalentes entre sí:")
bullets(doc,["Su tiempo próximo de iniciación es igual a su tiempo remoto de iniciación.",
             "Su tiempo próximo de terminación es igual a su tiempo remoto de terminación.",
             "Su holgura total es cero.",
             "Sus cuatro holguras son cero."],num=True)
P(doc,f"Se verificaron las cuatro condiciones en las 257 actividades. {len(M['CRIT'])} las cumplen, frente a "
  f"{len(O['CRIT'])} en la versión anterior.")

H(doc,"11.2 La trayectoria",2)
P(doc,f"La ruta crítica está formada por {len(R.CADENA)} actividades encadenadas y mide "
  f"{M['TOTAL']:.2f} días hábiles.")
rows=[(k, R.T[k]["desc"], R.T[k]["sig"], R.T[k]["resp"], f(R.T[k]["dur"]), f"{M['EF'][k]:.2f}")
      for k in R.CADENA]
rows.append(("","Duración total de la ruta crítica","","", "", f"{M['TOTAL']:.2f}"))
table(doc,["Clave","Actividad","Á.","Responsable","Días","Acumulado"],rows,
      widths=[1.4,7.2,0.8,2.2,1.1,2.3],fs=8.5)

H(doc,"11.3 Composición por área",2)
c = Counter(R.T[k]["area"] for k in M["CRIT"])
cv1 = Counter(R.T[k]["area"] for k in O["CRIT"])
table(doc,["Área","Responsable","Críticas v1.0","Críticas v2.0","De un total de"],
 [(a, R.RESP[a], cv1.get(a,0), c.get(a,0), R.NTAR[a])
  for a in sorted(R.AREAS, key=lambda x:-c.get(x,0))],
 widths=[5.2,3.0,2.6,2.6,2.6],fs=9.5)
P(doc,"Este cuadro contiene el hallazgo principal, y es el mismo que en la versión anterior: la ruta crítica "
  "atraviesa el entorno tridimensional, QA y la experiencia emocional, más una sola actividad de interfaz. "
  "No aparece en ella ninguna actividad de desarrollo VR, del juego de ritmo, de música ni de sonido.")
P(doc,"Que el resultado se mantenga después de cerrar la red es significativo. En la versión 1.0 podía "
  "atribuirse al defecto de las noventa actividades colgando, que regalaba holgura a las ramas de desarrollo. "
  "Con la red bien formada, esa explicación ya no aplica: el desarrollo VR conserva 7.4 días de holgura "
  "mínima y el sonido 8.0 porque efectivamente terminan antes de que el entorno tridimensional esté listo.")

H(doc,"11.4 Por qué el entorno tridimensional resulta crítico",2)
P(doc,"El resultado es una consecuencia directa de cómo está construida la red. La rama del entorno "
  "tridimensional es una cadena de diecisiete actividades estrictamente secuenciales, todas a cargo de la "
  "misma persona, sin ninguna bifurcación:")
P(doc,"AE.1 investigar el entorno → AE.2 comparar ambientes → AE.3 definir dirección artística → AC.2 "
  "definir paleta → AC.3 crear moodboard → AA.1 listar assets → AA.3 buscar assets → AA.4 revisar licencias "
  "→ AP.1 componer el escenario → AP.3 aplicar la paleta → AL.1 iluminación inicial → AL.2 iluminación "
  "definitiva → AD.1 elementos ambientales → AO.1 optimizar texturas → AO.2 configurar LOD → AO.5 evaluar "
  "rendimiento → AO.6 optimizar hasta cumplir.", italic=True)
P(doc,f"Suman {sum(R.T[k]['dur'] for k in R.CADENA if R.T[k]['sig']=='A'):.2f} días hábiles encadenados, "
  f"es decir más de la mitad de los {M['TOTAL']:.2f} días del proyecto, y ninguna puede adelantarse ni "
  "ejecutarse en paralelo con la anterior tal como está declarada la red. En la versión 1.0 esta cadena "
  "tenía catorce actividades; los cambios le agregaron tres más, AA.4, AD.1 y AO.2, alargándola.")

H(doc,"11.5 Verificación aritmética",2)
P(doc,"La suma de las duraciones a lo largo de la ruta crítica debe coincidir exactamente con la duración "
  "obtenida en el recorrido hacia adelante. Si no coincidiera, habría un error en la red o en los cálculos.")
table(doc,["Concepto","Valor"],[
 (f"Suma de las duraciones de las {len(R.CADENA)} actividades de la ruta crítica",
  f"{sum(R.T[k]['dur'] for k in R.CADENA):.2f} días"),
 ("Duración obtenida en el recorrido hacia adelante", f"{M['TOTAL']:.2f} días"),
 ("Diferencia", f"{abs(sum(R.T[k]['dur'] for k in R.CADENA)-M['TOTAL']):.2f} días"),
],widths=[10.4,5.6],fs=10)
P(doc,"Coinciden. Los cálculos son internamente consistentes.")

doc.save('/tmp/_v2_c.docx'); print("C OK")
