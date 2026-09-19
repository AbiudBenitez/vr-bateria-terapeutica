# -*- coding: utf-8 -*-
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from neutro import *
from docx import Document
from docx.shared import Cm
import rc257 as R, costos as K, compresion as CP
M = R.V2
def d(x): return f"${x:,.0f}"
def d2(x): return f"${x:,.2f}"
f = R.fnum
FIG="/Users/abiudbenitez/Documents/Claude/Projects/vr-bateria-terapeutica/entregables/figuras_costos/"
doc = Document('/tmp/_cost_b.docx')

# ================================================================ 9
H(doc,"9. Escenarios de sensibilidad",1)
P(doc,"El presupuesto valora el trabajo del equipo a tarifas de practicante ancladas al salario mínimo. Es "
  "una decisión metodológica que conviene explicitar, porque existen otras dos lecturas legítimas del costo "
  "de un proyecto académico.")
rows=[]
for nom,t,just in K.ESCENARIOS:
    rows.append((nom, d2(t) if t else "—", d(K.escenario(t)), just))
table(doc,["Escenario","Tarifa media","Costo del proyecto","Fundamento"],rows,
      widths=[4.4,2.2,2.6,6.8],fs=9.5)
P(doc,"El escenario B es la línea base de este documento. El escenario A es el desembolso real del equipo: "
  "el trabajo se cursa por créditos y no se remunera, de modo que no sale dinero de ningún bolsillo. El "
  "escenario C dimensiona lo que costaría ejecutar el mismo alcance con personal titulado.")
P(doc,"Presentar los tres tiene un propósito: distinguir el costo del valor. El proyecto no cuesta dinero, "
  "pero el trabajo que se invierte en él sí tiene un valor, y ese valor es el que se administra.")

# ================================================================ 10
H(doc,"10. Tabla de simultaneidad y compresión de la red",1)
H(doc,"10.1 Trabajo simultáneo",2)
P(doc,"La tabla indica qué trabajo puede ejecutarse al mismo tiempo. Se deriva directamente de la red de "
  "precedencias: dos actividades pueden ser simultáneas cuando sus intervalos entre tiempo próximo de "
  "iniciación y de terminación se traslapan y ninguna depende de la otra.")
rows=[]
for a,b in CP.TRAMOS:
    act = CP.activas(a,b)
    if not act: continue
    tot = sum(len(v) for v in act.values())
    det = " · ".join(f"{s}: {len(v)}" for s,v in sorted(act.items()))
    rows.append((f"{a:.0f} – {b:.0f}", len(act), tot, det))
table(doc,["Días","Áreas activas","Tareas en curso","Detalle por área"],rows,
      widths=[2.0,2.2,2.4,9.4],fs=9.5)
P(doc,"Las claves de área son D desarrollo VR, J juego de ritmo, E experiencia emocional, M música, "
  "S sonido, I interfaz, A entorno tridimensional y Q aseguramiento de calidad.",italic=True)
P(doc,"El proyecto alcanza su máximo paralelismo en los primeros diez días hábiles, con las ocho áreas "
  "activas a la vez. Hacia el final solo queda activa el área de QA, lo que explica tanto la forma de la "
  "curva S como la sobrecarga de esa área.")

H(doc,"10.2 Compresión de la red",2)
P(doc,"La compresión responde a una pregunta económica: si hiciera falta terminar antes, ¿cuánto costaría "
  "cada día ganado? Solo tiene sentido comprimir actividades de la ruta crítica.")
P(doc,"Cada área tiene una sola persona asignada, de modo que no es posible comprimir añadiendo un segundo "
  "recurso. La única vía es extender la jornada: de ocho a diez horas, con las dos horas adicionales "
  "retribuidas al 150 %, lo que permite reducir la duración hasta un 20 %.")
pts,_ = CP.curva()
table(doc,["Punto","Duración","Sobrecosto acumulado","Costo por día ganado"],[
 ("Duración normal", f"{pts[0][0]:.2f} días", "Sin sobrecosto", "—"),
 ("Compresión máxima", f"{pts[-1][0]:.2f} días", d(pts[-1][1]), d(pts[-1][1]/(pts[0][0]-pts[-1][0]))),
],widths=[4.4,3.0,4.2,4.4],fs=9.5)
if os.path.exists(FIG+"Fig2_Curva_compresion.png"):
    doc.add_picture(FIG+"Fig2_Curva_compresion.png", width=Cm(15.4)); doc.paragraphs[-1].alignment=C
caption(doc,"Fig. 2 — Curva de compresión tiempo-costo.")
P(doc,f"El proyecto puede reducirse de {pts[0][0]:.2f} a {pts[-1][0]:.2f} días hábiles, es decir "
  f"{pts[0][0]-pts[-1][0]:.2f} días, con un sobrecosto de {d(pts[-1][1])}. No lo necesita para cumplir el "
  f"calendario: la ruta crítica mide {pts[0][0]:.2f} días y hay {K.DIAS_DISPONIBLES} disponibles. La "
  "compresión es un instrumento de contingencia.")
P(doc,"Obsérvese que la pendiente de costo resulta idéntica dentro de cada perfil. Bajo el modelo de "
  "jornada extendida, la pendiente es exactamente cuatro veces la tarifa horaria, con independencia de "
  "cuánto dure la actividad. La conclusión práctica: conviene comprimir primero el trabajo de los perfiles "
  "de menor nivel que estén sobre la ruta crítica.")
P(doc,"Su límite conviene declararlo: la compresión resuelve problemas de ruta crítica, no de carga de "
  "trabajo. La restricción dominante del proyecto es que el área de QA tiene 511 horas asignadas y su "
  "ventana permite unas 306. Comprimir no corrige eso; redistribuir trabajo, sí.")

doc.add_page_break()
doc.save('/tmp/_cost_c.docx'); print("C OK")
