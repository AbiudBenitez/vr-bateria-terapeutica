# -*- coding: utf-8 -*-
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from neutro import *
from docx import Document
import rc257 as R, costos as K, pruebas as PR
doc = Document('/tmp/_pcr_d.docx')

doc.add_page_break()
H(doc,"24. Implementar la respuesta a los riesgos",2)
P(doc,"Planificar respuestas no reduce ningún riesgo: la reducción ocurre cuando alguien las ejecuta. Es "
  "donde la mayoría de los registros de riesgos fracasa, porque quedan escritos y nadie los activa.")
H(doc,"24.1 Regla de activación",3)
P(doc,"Cuando el responsable observa el disparador, ejecuta la respuesta sin esperar autorización ni a la "
  "reunión semanal. Solo se escala cuando la respuesta implica recortar alcance o consumir más de tres días "
  "de reserva.")
table(doc,["Situación","Quién decide","Plazo"],[
 ("El disparador se cumple y la respuesta está dentro del área del responsable.","El responsable del riesgo, por su cuenta.","De inmediato. Se informa en la siguiente reunión."),
 ("La respuesta requiere apoyo de otra área.","El gerente de proyecto.","Dentro de las 24 horas siguientes."),
 ("La respuesta implica recortar alcance o consumir más de tres días de reserva.","El director de proyecto.","En la siguiente reunión, o antes si es urgente."),
 ("Aparece un riesgo no registrado con impacto mayor a dos días.","El gerente lo registra y asigna responsable.","El mismo día en que se detecta."),
],widths=[6.8,4.8,4.4],fs=9.5)
H(doc,"24.2 Registro de ejecución",3)
P(doc,"Cada vez que se ejecuta una respuesta se anota qué ocurrió. Sin ese registro no se puede saber si la "
  "reserva sigue alcanzando ni aprender nada para el siguiente proyecto.")
table(doc,["Campo","Contenido"],[
 ("Fecha","Cuándo se observó el disparador."),
 ("Riesgo","Identificador del registro."),
 ("Disparador observado","El hecho concreto, no la impresión."),
 ("Respuesta ejecutada","Qué se hizo, y quién."),
 ("Días de reserva consumidos","Cuánto se gastó de los "+f"{K.RESERVA_CRONO:.2f}"+" disponibles."),
 ("Reserva restante","Saldo después de la ejecución."),
 ("Resultado","Si la respuesta funcionó y qué quedó pendiente."),
],widths=[4.4,11.6],fs=9.5)

H(doc,"25. Monitorear los riesgos",2)
table(doc,["Actividad","Frecuencia","Responsable","Producto"],[
 ("Revisión del registro","Semanal","Gerente de proyecto","Registro con probabilidades reevaluadas"),
 ("Verificación de disparadores","Semanal","Responsable de cada riesgo","Confirmación de que ninguno se cumplió, o activación de la respuesta"),
 ("Identificación de riesgos nuevos","Semanal","Todo el equipo","Altas en el registro, con responsable el mismo día"),
 ("Control del saldo de la reserva","Semanal","Gerente de proyecto","Días consumidos y días restantes"),
 ("Reevaluación completa","En cada hito de control","Gerente y director","Registro revisado y reserva recalculada"),
 ("Cierre de riesgos superados","En cada hito","Gerente de proyecto","Riesgos cerrados con su lección aprendida"),
 ("Auditoría de riesgos","A la mitad del proyecto","Gerente de proyecto","Informe sobre la eficacia de las respuestas"),
],widths=[4.6,2.4,3.4,5.6],fs=9.5)
H(doc,"25.1 Cierre de riesgos",3)
P(doc,"Un riesgo se cierra cuando la ventana en que podía materializarse ya pasó. Cerrarlo libera la parte "
  "de reserva que tenía asociada.")
table(doc,["Riesgo","Se cierra cuando","Libera"],[
 ("R1 Latencia sobre el umbral","Se supera la medición del caso CN-14 con resultado favorable.","1.50 días"),
 ("R5 Préstamo de visores","Existe confirmación escrita de la Facultad.","1.20 días"),
 ("R13 Actualización del motor","La versión queda fijada y verificada en todos los equipos.","0.60 días"),
 ("R18 Usuarios para pruebas","Se confirman cinco participantes.","0.60 días"),
 ("R16 Choque con exámenes","Concluye el periodo de parciales.","1.65 días"),
],widths=[4.4,8.0,3.6],fs=9.5)
H(doc,"25.2 Indicadores de seguimiento",3)
table(doc,["Indicador","Qué señala","Umbral de alarma"],[
 ("Días de reserva consumidos","Cuánto margen queda.","Superar el 50 % antes de la mitad del proyecto."),
 ("Riesgos nuevos por semana","Si la identificación inicial fue suficiente.","Más de dos semanas seguidas con altas."),
 ("Riesgos con disparador cumplido y respuesta sin ejecutar","Si el plan se aplica o solo está escrito.","Cualquier valor distinto de cero."),
 ("Valor esperado total del registro","Si la exposición crece o baja.","Superar la suma de reserva más compresión disponible."),
 ("Áreas sin reportar avance","Si el control del cronograma está ciego.","Cualquier área con dos semanas sin reportar."),
],widths=[4.6,6.4,5.0],fs=9.5)
P(doc,"El tercero es el indicador decisivo. Un registro con disparadores cumplidos y respuestas sin ejecutar "
  "no es una herramienta de gestión: es documentación.")

doc.add_page_break()
H(doc,"26. Resumen ejecutivo",1)
table(doc,["Concepto","Valor"],[
 ("Presupuesto del proyecto", f"${K.PRESUPUESTO:,.0f}"),
 ("Esfuerzo estimado", f"{K.HORAS_TOT:,.0f} horas sobre 257 actividades"),
 ("Tarifa media", f"${K.TARIFA_MED:,.2f} por hora, {K.TARIFA_MED/K.SM_HORA:.2f} salarios mínimos"),
 ("Costo de la calidad", f"${sum(K.COQ_TOT.values()):,.0f}, {100*sum(K.COQ_TOT.values())/K.MANO_OBRA:.1f} % de la mano de obra"),
 ("Casos de prueba diseñados", f"{len(PR.CAJA_BLANCA)} de caja blanca y {len(PR.CAJA_NEGRA)} de caja negra"),
 ("Riesgos identificados", f"{len(PR.RIESGOS)} en siete categorías"),
 ("Valor esperado del riesgo", f"{PR.EMV_RESERVA:.2f} días hábiles que consumen reserva"),
 ("Reserva de cronograma", f"{K.RESERVA_CRONO:.2f} días hábiles, el {100*K.RESERVA_CRONO/PR.EMV_RESERVA:.0f} % de la exposición"),
],widths=[7.0,9.0],fs=10)
H(doc,"26.1 Puntos que requieren decisión del equipo",2)
table(doc,["#","Punto","Por qué importa"],[
 ("1","La reserva de cronograma cubre dos tercios de la exposición al riesgo.",
  "Comprimir la red aporta 4.10 días más por $901 y sube la cobertura al 93 %. Si no basta, hay que recortar alcance secundario."),
 ("2","Nadie registra las horas efectivamente trabajadas.",
  "Sin ese dato no se pueden calcular los índices de costo y la mitad del método de valor ganado queda inutilizable."),
 ("3","Las pruebas de caja blanca las ejecuta quien programó el módulo.",
  "Es correcto y necesario, pero exige que cada quien reserve tiempo para probar su propio código, no solo para escribirlo."),
 ("4","Confirmar por escrito el préstamo de los visores.",
  "R5 es el riesgo de mayor impacto unitario del registro, con seis días hábiles."),
 ("5","Redistribuir la documentación del área de QA.",
  "Es la restricción dominante del proyecto y el riesgo R3, con el tercer valor esperado más alto."),
 ("6","Designar un segundo conocedor por área.",
  "Es la respuesta a R14. Hoy, si alguien se da de baja, se pierde el conocimiento de su área completa."),
],widths=[0.8,6.4,8.8],fs=9.5)

out="/Users/abiudbenitez/Documents/Claude/Projects/vr-bateria-terapeutica/entregables/Planes_Costos_Calidad_Riesgos.docx"
doc.save(out); print("OK", out)
print("párrafos:",len(doc.paragraphs)," tablas:",len(doc.tables)," imágenes:",len(doc.inline_shapes))
