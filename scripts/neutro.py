# -*- coding: utf-8 -*-
"""Formato neutral: estilos por defecto de Word, sin colores ni sombreados."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH as WA
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

C = WA.CENTER
J = WA.JUSTIFY

def nuevo(margen=2.5, fuente=None, tam=None):
    """Documento en blanco. Si se indica fuente o tamaño, se aplican a todo el texto."""
    doc = Document()
    s = doc.sections[0]
    s.left_margin=s.right_margin=Cm(margen)
    s.top_margin=s.bottom_margin=Cm(2.5)
    if fuente or tam: tipografia(doc, fuente, tam)
    return doc

def tipografia(doc, fuente=None, tam=None):
    """Fija la fuente base del documento y la de los encabezados.

    Word toma la fuente de tres lugares distintos: la del estilo, la de Asia oriental
    y la de escritura compleja. Si solo se cambia la primera, algunos caracteres se
    siguen componiendo con la fuente anterior, así que se fijan las tres."""
    def aplicar(st, sz=None):
        if fuente:
            st.font.name = fuente
            rpr = st.element.get_or_add_rPr()
            rf = rpr.find(qn('w:rFonts'))
            if rf is None:
                rf = OxmlElement('w:rFonts'); rpr.append(rf)
            for at in ('w:ascii','w:hAnsi','w:eastAsia','w:cs'):
                rf.set(qn(at), fuente)
        if sz: st.font.size = Pt(sz)
    aplicar(doc.styles['Normal'], tam)
    if tam:
        for h, f in (('Heading 1',1.45), ('Heading 2',1.20), ('Heading 3',1.05)):
            try: aplicar(doc.styles[h], round(tam*f,1))
            except KeyError: pass
    elif fuente:
        for h in ('Heading 1','Heading 2','Heading 3'):
            try: aplicar(doc.styles[h])
            except KeyError: pass
    for est in ('List Bullet','List Number','Table Grid'):
        try: aplicar(doc.styles[est])
        except KeyError: pass
    return doc

def campo(doc, etiqueta, valor, after=4, tam=None):
    """Renglón de portada con la etiqueta en negrita y el valor en texto normal."""
    p = doc.add_paragraph()
    r = p.add_run(etiqueta + ": "); r.bold = True
    v = p.add_run(valor)
    if tam:
        r.font.size = Pt(tam); v.font.size = Pt(tam)
    p.paragraph_format.space_after = Pt(after)
    return p

def P(doc, txt="", bold=False, italic=False, size=None, align=None, after=8):
    p=doc.add_paragraph(); r=p.add_run(txt); r.bold=bold; r.italic=italic
    if size: r.font.size=Pt(size)
    p.alignment = align if align else J
    p.paragraph_format.space_after=Pt(after)
    return p

def H(doc, txt, lvl=1):
    return doc.add_heading(txt, level=lvl)

def bullets(doc, items, num=False):
    st = 'List Number' if num else 'List Bullet'
    for it in items:
        p=doc.add_paragraph(it, style=st)
        p.paragraph_format.space_after=Pt(3)

def table(doc, headers, rows, widths=None, fs=9.0, negritas_col0=False, autofit=True):
    t=doc.add_table(rows=1, cols=len(headers))
    t.style='Table Grid'
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=autofit
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=""
        pr=c.paragraphs[0]; pr.alignment=C
        pr.paragraph_format.space_after=Pt(2)
        r=pr.add_run(str(h)); r.bold=True; r.font.size=Pt(fs)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=""
            pr=cells[i].paragraphs[0]
            pr.paragraph_format.space_after=Pt(2)
            r=pr.add_run(str(v)); r.font.size=Pt(fs)
            if negritas_col0 and i==0: r.bold=True
    if widths:
        t.autofit=False
        for r_ in t.rows:
            for i,w in enumerate(widths): r_.cells[i].width=Cm(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t

def caption(doc, txt):
    p=doc.add_paragraph(); p.alignment=C
    r=p.add_run(txt); r.italic=True; r.font.size=Pt(9)
    p.paragraph_format.space_after=Pt(10)

def portada(doc, titulo, subtitulo, datos, sub2=None):
    for _ in range(3): doc.add_paragraph()
    P(doc,"UNIVERSIDAD AUTÓNOMA DE NUEVO LEÓN",bold=True,size=12,align=C,after=2)
    P(doc,"FACULTAD DE INGENIERÍA MECÁNICA Y ELÉCTRICA",bold=True,size=11,align=C,after=30)
    P(doc,titulo,bold=True,size=18,align=C,after=6)
    P(doc,subtitulo,size=12,align=C,after=4)
    if sub2: P(doc,sub2,italic=True,size=10.5,align=C,after=24)
    else: doc.add_paragraph().paragraph_format.space_after=Pt(24)
    table(doc,["Campo","Dato"],datos,widths=[4.5,11.0],fs=10)
    doc.add_page_break()

def landscape(doc, margen=1.5):
    s=doc.add_section(WD_SECTION.NEW_PAGE)
    s.orientation=WD_ORIENT.LANDSCAPE
    w,h=s.page_width,s.page_height
    if w<h: s.page_width,s.page_height=h,w
    s.left_margin=s.right_margin=s.top_margin=s.bottom_margin=Cm(margen)
    return s

def portrait(doc, margen=2.5):
    s=doc.add_section(WD_SECTION.NEW_PAGE)
    s.orientation=WD_ORIENT.PORTRAIT
    w,h=s.page_width,s.page_height
    if w>h: s.page_width,s.page_height=h,w
    s.left_margin=s.right_margin=Cm(margen); s.top_margin=s.bottom_margin=Cm(2.5)
    return s
