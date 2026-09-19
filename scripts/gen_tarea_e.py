# -*- coding: utf-8 -*-
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from neutro import *
from docx import Document
import pmbok as PB
doc = Document('/tmp/_tarea_d.docx')

# ================================================================ 6
H(doc,"6. Conclusión",1)
P(doc,"La Guía del PMBOK organiza la dirección de proyectos cruzando dos clasificaciones independientes: "
  "cinco grupos de procesos que responden a cuándo se hace el trabajo, y diez áreas de conocimiento que "
  "responden a qué se gestiona. De ese cruce resultan 49 procesos, cada uno descrito mediante entradas, "
  "herramientas y salidas.")
P(doc,"La tabla de la sección 3.3 cumple el objetivo planteado: muestra que cada grupo recibe información "
  "de los anteriores y produce documentos que alimentan a los siguientes. El acta de constitución nace en "
  "el inicio y se convierte en entrada de casi toda la planificación; las líneas base nacen en la "
  "planificación y se vuelven la referencia contra la que el monitoreo mide el desempeño; los entregables "
  "nacen en la ejecución y el cierre los transfiere formalmente. Ese encadenamiento, y no la lista de "
  "procesos, es lo que constituye el método.")
P(doc,"El contraste con el proyecto de la unidad de aprendizaje confirmó que la estructura es aplicable y, "
  f"a la vez, reveló una omisión concreta: de los 49 procesos se han ejecutado 23, con buena cobertura en "
  "planificación pero sin ningún proceso del área de comunicaciones. En un equipo de ocho personas "
  "distribuidas en áreas distintas, esa carencia tiene consecuencias prácticas.")
P(doc,"Conviene cerrar con una precisión sobre la naturaleza del estándar. El PMBOK no prescribe seguir los "
  "49 procesos en todo proyecto: establece que deben adaptarse a las condiciones de cada caso. El propio "
  "proyecto analizado lo ilustra, al no ejecutar ningún proceso de adquisiciones porque no realiza compras. "
  "Distinguir entre omitir por adaptación y omitir por descuido es, probablemente, la habilidad que el "
  "estándar exige de quien lo aplica.")

# ================================================================ 7
H(doc,"7. Lo aprendido",1)
P(doc,"Antes de esta actividad yo entendía el PMBOK como una lista de cosas que hay que documentar. Lo veía "
  "como un requisito de la materia más que como una herramienta. Construir la tabla de entradas y salidas "
  "cambió esa impresión, y lo hizo por una razón concreta: al tener que escribir qué recibe y qué produce "
  "cada grupo, quedó a la vista que los documentos no son independientes entre sí.")
P(doc,"El momento en que eso se volvió evidente fue al revisar nuestro propio proyecto. Habíamos elaborado "
  "el acta constitutiva, después la estructura de desglose, después el cronograma y después el presupuesto, "
  "y siempre los tratamos como entregables separados que había que ir entregando. Al ubicarlos en la matriz "
  "me di cuenta de que forman una cadena: el presupuesto se calculó sobre las actividades del cronograma, "
  "que salieron de la estructura de desglose, que salió del alcance del acta. Si el acta hubiera estado mal, "
  "todo lo demás habría heredado el error.")
P(doc,"De hecho eso nos pasó. El presupuesto del proyecto se recalculó tres veces, y ninguna corrección fue "
  "un error de aritmética: las tres veces cambió un supuesto de entrada. Primero las tarifas, luego el tipo "
  "de recurso, luego la figura de contratación. Ahora entiendo que el modelo de entradas y salidas está "
  "describiendo exactamente ese fenómeno, y que revisar las entradas antes de ejecutar un proceso ahorra "
  "rehacer lo que viene después.")
P(doc,"Lo segundo que aprendí es a leer la distribución de los procesos como una postura y no como un dato. "
  "Que 24 de 49 procesos estén en planificación y solo 10 en ejecución me pareció desproporcionado al "
  "principio. Después, viendo que en nuestro proyecto los problemas serios no fueron técnicos sino de "
  "planificación (la red de dependencias que no cerraba, la sobrecarga de un área, el presupuesto mal "
  "estimado), la proporción empezó a tener sentido.")
P(doc,"Lo tercero es más incómodo de admitir. Al mapear nuestros documentos contra los 49 procesos apareció "
  "que no tenemos plan de comunicaciones. No lo habíamos notado porque nadie lo pidió como entregable. Eso "
  "me hizo ver para qué sirve realmente un estándar: no para tener qué entregar, sino para darse cuenta de "
  "lo que falta cuando nadie lo está preguntando.")
P(doc,"Por último, revisar la evolución hacia la séptima y la octava edición me dejó claro que el estándar "
  "también se equivoca y se corrige. La séptima abandonó los procesos y la octava los reincorporó cinco años "
  "después. Saber que un cuerpo de conocimiento profesional cambia de opinión hace más fácil usarlo con "
  "criterio en vez de aplicarlo al pie de la letra.")

# ================================================================ 8
doc.add_page_break()
H(doc,"8. Bibliografía",1)
P(doc,"Referencias en formato APA, séptima edición.")
REFS = [
 "Project Management Institute. (2017). Guía de los fundamentos para la dirección de proyectos "
 "(Guía del PMBOK) (6.ª ed.). Project Management Institute.",
 "Project Management Institute. (2021). Guía de los fundamentos para la dirección de proyectos "
 "(Guía del PMBOK) y El estándar para la dirección de proyectos (7.ª ed.). Project Management Institute.",
 "Project Management Institute. (2026). A guide to the project management body of knowledge "
 "(PMBOK guide) and The standard for project management (8.ª ed.). Project Management Institute.",
 "Project Management Institute. (2026). PMBOK Guide. https://www.pmi.org/standards/pmbok",
 "Neira Tovar, L. A. (2026). Administración de proyectos, parte IV [Presentación de clase]. "
 "Facultad de Ingeniería Mecánica y Eléctrica, Universidad Autónoma de Nuevo León.",
 "Montaño, A. (s. f.). El método de la ruta crítica [Documento de apoyo del curso]. "
 "Facultad de Ingeniería Mecánica y Eléctrica, Universidad Autónoma de Nuevo León.",
 "Equipo A. (2026). Acta constitutiva del proyecto: simulación de batería en realidad virtual con "
 "juego de ritmo, versión 3.0 [Documento de proyecto]. Universidad Autónoma de Nuevo León.",
 "Equipo A. (2026). Ruta crítica del proyecto: simulación VR de batería con juego de ritmo, "
 "versión 3.0 [Documento de proyecto]. Universidad Autónoma de Nuevo León.",
 "Equipo A. (2026). Análisis de costos, calidad y riesgos [Documento de proyecto]. "
 "Universidad Autónoma de Nuevo León.",
 "Comisión Nacional de los Salarios Mínimos. (2026). Salarios mínimos generales y profesionales "
 "vigentes a partir del 1 de enero de 2026. Gobierno de México. https://www.gob.mx/conasami",
]
for r in REFS:
    p = doc.add_paragraph(r)
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.first_line_indent = Cm(-1.0)
    p.paragraph_format.space_after = Pt(8)
    for run in p.runs: run.font.size = Pt(10)

P(doc,"Nota sobre las fuentes: las ediciones séptima y octava de la guía se consultaron a través del sitio "
  "oficial del Project Management Institute para verificar su estructura y su fecha de publicación. La "
  "documentación del proyecto citada es de elaboración propia del equipo durante el semestre y se incluye "
  "porque constituye la evidencia de la sección 5.2.", italic=True, size=9.5)

out="/Users/abiudbenitez/Documents/Claude/Projects/vr-bateria-terapeutica/tareas-individuales/2026-09-14-tabla-pmbok/Elementos_Estudio_PMBOK.docx"
os.makedirs(os.path.dirname(out), exist_ok=True)
doc.save(out); print("OK", out)
print("párrafos:",len(doc.paragraphs)," tablas:",len(doc.tables)," imágenes:",len(doc.inline_shapes))
