# -*- coding: utf-8 -*-
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from neutro import *
from docx import Document
from docx.shared import Cm
import rc257 as R, costos as K
M = R.V2
def d(x): return f"${x:,.0f}"
def d2(x): return f"${x:,.2f}"
FIG="/Users/abiudbenitez/Documents/Claude/Projects/vr-bateria-terapeutica/entregables/figuras-costos/"
doc = Document('/tmp/_cost_a.docx')

# ================================================================ 6
H(doc,"6. Asignación de perfil a cada grupo de actividades",1)
P(doc,"La estimación ascendente exige asignar a cada actividad el recurso que efectivamente la ejecuta. Se "
  "asignó por grupo de claves y no por área, porque dentro de un área conviven trabajos de perfiles "
  "distintos. El caso más claro es el área de QA, documentación y gestión, que reúne planificación del "
  "proyecto, pruebas y redacción de manuales: cobrar las tres al mismo nivel distorsionaría el resultado.")
rows=[]
for s in "DJEMSIAQ":
    gs = sorted({g for (sig,g) in K.POR_GRUPO if sig==s})
    for n,g in enumerate(gs):
        cnt,h,c = K.POR_GRUPO[(s,g)]
        p = K.GRUPO[g]
        rows.append((s if n==0 else "", g, p, f"{K.multiplo(p):.1f}", d2(K.PERFIL[p][0]), cnt, f"{h:,.0f}", d(c)))
table(doc,["Á.","Grupo","Perfil asignado","SM","$/hora","Tareas","Horas","Costo"],rows,
      widths=[0.7,1.2,4.4,0.9,1.5,1.2,1.3,2.0],fs=7.5)
P(doc,"Son 68 grupos. La asignación completa por actividad individual está en el anexo A.")

# ================================================================ 7
H(doc,"7. Agregación ascendente del costo",1)
P(doc,"El costo sube por tres niveles: de la actividad al grupo de claves, del grupo al área y del área al "
  "proyecto. Cada nivel es una cuenta de control sobre la que se mide el desempeño.")
H(doc,"7.1 Por área de trabajo",2)
rows=[]
for a in sorted(R.AREAS, key=lambda x:-K.POR_AREA[x][2]):
    n,h,c = K.POR_AREA[a]
    rows.append((R.SIGLA[a], a, R.RESP[a], n, f"{h:,.0f}", d(c), f"{100*c/K.MANO_OBRA:.1f} %"))
rows.append(("","Total de mano de obra","", sum(K.POR_AREA[a][0] for a in R.AREAS),
             f"{K.HORAS_TOT:,.0f}", d(K.MANO_OBRA), "100.0 %"))
table(doc,["Á.","Área","Responsable","Tareas","Horas","Costo","%"],rows,
      widths=[0.8,4.4,2.6,1.4,1.6,2.4,1.4],fs=9.5)
P(doc,"El área de QA, documentación y gestión concentra la cuarta parte del costo de mano de obra. No es "
  "un error de estimación: reúne 60 de las 257 actividades y, además de las pruebas, absorbe la "
  "planificación, el seguimiento y toda la documentación del proyecto, porque la organización del equipo no "
  "asignó a nadie los roles de dirección, gerencia y coordinación.")

H(doc,"7.2 Por perfil profesional",2)
rows=[]
for p in sorted(K.POR_PERFIL, key=lambda x:-K.POR_PERFIL[x][2]):
    n,h,c = K.POR_PERFIL[p]
    rows.append((p, f"{K.multiplo(p):.1f} SM", d2(K.PERFIL[p][0]), n, f"{h:,.0f}", d(c), f"{100*c/K.MANO_OBRA:.1f} %"))
table(doc,["Perfil","Múltiplo","$/hora","Tareas","Horas","Costo","%"],rows,
      widths=[5.0,1.6,1.6,1.4,1.6,2.2,1.4],fs=9.5)

H(doc,"7.3 Distribución del costo",2)
if os.path.exists(FIG+"Fig3_Distribucion_costo.png"):
    doc.add_picture(FIG+"Fig3_Distribucion_costo.png", width=Cm(16.4)); doc.paragraphs[-1].alignment=C
caption(doc,"Fig. 3 — Distribución del costo de mano de obra por área y por perfil.")

doc.add_page_break()
# ================================================================ 8
H(doc,"8. Determinar el presupuesto",1)
H(doc,"8.1 El proyecto no realiza compras",2)
P(doc,"El presupuesto de este proyecto se compone únicamente de mano de obra. No hay partidas de equipo, "
  "licencias ni operación, y conviene justificar cada exclusión para que no se lea como una omisión.")
table(doc,["Recurso","Descripción","Condición"],
 [(n,des,cond) for n,des,cond in K.PROPORCIONADOS],
 widths=[4.2,7.0,4.8],fs=9.5)
P(doc,"Conforme a la guía PMBOK, los recursos que aporta la organización ejecutante no se cargan al "
  "presupuesto del proyecto, pero deben quedar registrados: su disponibilidad es un supuesto del que el "
  "proyecto depende. El préstamo de los visores genera el riesgo R5 de la sección 12.")
P(doc,"Las partidas siguientes existirían en un proyecto comercial y aquí se resuelven sin costo:")
table(doc,["Partida","Cómo se resuelve sin costo"],
 [(n,c) for n,c in K.SIN_COSTO], widths=[5.6,10.4],fs=9.5)

H(doc,"8.2 Reservas",2)
P(doc,"La guía PMBOK distingue dos reservas. La de contingencia cubre el impacto de riesgos identificados y "
  "forma parte de la línea base; la de gestión cubre trabajo imprevisto no identificado y queda fuera de "
  "ella. Ambas son monetarias por definición.")
P(doc,"En este proyecto ambas son cero, y la razón es estructural, no una omisión. La reserva de "
  "contingencia existe para financiar el impacto económico de los riesgos. Como el proyecto no compra nada, "
  "no contrata a nadie y no paga penalizaciones, ninguno de los nueve riesgos identificados tiene impacto "
  "monetario: todos impactan el cronograma o el alcance. Financiar con dinero un riesgo que no cuesta dinero "
  "no tiene sentido.")
P(doc,"En su lugar se establece una reserva de cronograma, que es la respuesta adecuada a riesgos cuyo "
  "efecto se mide en días:")
table(doc,["Concepto","Días hábiles"],[
 ("Disponibles del 7 de septiembre al 13 de noviembre de 2026", f"{K.DIAS_DISPONIBLES}"),
 ("Duración de la red de actividades", f"{K.DURACION_RED:.2f}"),
 ("Reserva de cronograma", f"{K.RESERVA_CRONO:.2f}"),
 ("Valor esperado del impacto de los riesgos, sección 12", f"{K.EMV_DIAS:.2f}"),
 ("Margen de la reserva sobre el valor esperado", f"{K.RESERVA_CRONO-K.EMV_DIAS:.2f}"),
],widths=[11.0,5.0],fs=10)
P(doc,f"La reserva de cronograma es suficiente: {K.RESERVA_CRONO:.2f} días contra un valor esperado de "
  f"{K.EMV_DIAS:.2f}. El margen es de {K.RESERVA_CRONO-K.EMV_DIAS:.2f} días, es decir un "
  f"{100*(K.RESERVA_CRONO/K.EMV_DIAS-1):.0f} % por encima de lo que el análisis de riesgos exige. No es "
  "holgado, pero alcanza.")

H(doc,"8.3 Integración del presupuesto",2)
table(doc,["Concepto","Monto"],[
 ("Mano de obra", d(K.MANO_OBRA)),
 ("Costos no laborales", "Ninguno"),
 ("Costos directos", d(K.DIRECTOS)),
 ("Reserva de contingencia monetaria", "Ninguna. Se sustituye por reserva de cronograma"),
 ("Línea base de costos", d(K.LINEA_BASE)),
 ("Reserva de gestión", "Ninguna"),
 ("Presupuesto hasta la conclusión", d(K.PRESUPUESTO)),
],widths=[8.0,8.0],fs=10)

H(doc,"8.4 Curva S",2)
P(doc,"La curva S muestra cómo se acumula el valor planificado a lo largo del proyecto. Es la referencia "
  "contra la que se compara el avance real durante la ejecución.")
if os.path.exists(FIG+"Fig1_Curva_S.png"):
    doc.add_picture(FIG+"Fig1_Curva_S.png", width=Cm(16.4)); doc.paragraphs[-1].alignment=C
caption(doc,"Fig. 1 — Curva S del valor planificado, con el corte de medio curso.")
P(doc,"La curva es marcadamente frontal: cerca del 60 % del valor se acumula en la primera mitad del "
  "proyecto. La razón está en la red de precedencias: las ocho áreas arrancan simultáneamente el primer día "
  "y el trabajo se concentra en las primeras cuatro semanas, mientras que la última parte corresponde casi "
  "por completo al área de QA.")

doc.add_page_break()
doc.save('/tmp/_cost_b.docx'); print("B OK")
