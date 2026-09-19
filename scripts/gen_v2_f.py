# -*- coding: utf-8 -*-
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from neutro import *
from docx import Document
import rc257 as R
M, O = R.V2, R.ORIG
f = R.fnum
doc = Document('/tmp/_v2_e.docx')

doc.add_page_break()
# ================================================================ 15
H(doc,"15. Conclusiones",1)

H(doc,"15.1 Las tres cifras y qué mide cada una",2)
table(doc,["Cifra","Qué mide","Cuándo es la cifra correcta"],[
 (f"{O['TOTAL']:.2f} días","Trayectoria más larga de la red tal como estaba declarada antes del 9 de septiembre.",
  "Ya no. Se conserva solo para el contraste de la sección 12."),
 (f"{M['TOTAL']:.2f} días","Trayectoria más larga con la red cerrada.",
  "Si el equipo tuviera personal ilimitado. Es la duración correcta de la ruta crítica."),
 (f"{max(R.CARGA.values()):.1f} días","Tiempo que tarda la persona más cargada en ejecutar su trabajo.",
  "Con el reparto actual de ocho personas, una por área. Es la cifra que debe usarse para planear."),
],widths=[2.4,6.6,7.0],fs=9.5)

H(doc,"15.2 Hallazgos",2)
bullets(doc,[
 "Los cambios de dependencias corrigen el defecto de fondo señalado en el análisis anterior. Las actividades "
 "sin consecuente pasan de noventa a una, y las 257 actividades alcanzan ahora el entregable final. La red "
 "está bien formada.",
 f"La duración de la ruta crítica sube de {O['TOTAL']:.2f} a {M['TOTAL']:.2f} días hábiles. El aumento no es "
 "un empeoramiento del proyecto: es la duración que siempre tuvo, que la red anterior ocultaba.",
 "La ruta crítica sigue sin atravesar el desarrollo VR ni el juego de ritmo. Ahora es un resultado más "
 "sólido, porque ya no puede atribuirse a las actividades que colgaban.",
 "La cadena del entorno tridimensional creció de catorce a diecisiete actividades seriales y suma 19.25 de "
 "los 38.25 días del proyecto. Más de la mitad de la duración depende de una sola persona trabajando en "
 "secuencia.",
 f"El margen de calendario es de {49-M['TOTAL']:.2f} días hábiles sobre los 49 disponibles del 7 de "
 "septiembre al 13 de noviembre.",
 f"{sum(1 for a in R.AREAS if R.CARGA[a]>R.VENTANA[a])} de las ocho áreas están sobreasignadas. QA requiere "
 f"trabajar al {100*R.CARGA['QA, documentación y gestión']/R.VENTANA['QA, documentación y gestión']:.0f} % "
 "de su jornada disponible.",
 "Medido por carga de trabajo, el proyecto necesita 63.9 días hábiles y hay 47 disponibles antes del inicio "
 "de exámenes.",
])

H(doc,"15.3 Lo que corresponde decidir al equipo",2)
table(doc,["#","Punto a decidir","Por qué importa"],[
 ("1","Cómo se reparte la carga de QA, documentación y gestión.","Es la restricción que impide que el proyecto quepa en el calendario. Concentra 63.9 de los 268.2 días de carga total."),
 ("2","Quién asume dirección, gerencia y coordinación.","Hoy nadie los tiene asignados y su carga cayó dentro de QA."),
 ("3","Si las pruebas con usuarios deben esperar a la batería corregida.","DR.2 no es predecesora de QT.8. Agregarla no mueve la fecha final, de modo que la decisión es solo de criterio."),
 ("4","Si alguna de las diecisiete actividades del entorno tridimensional puede traslaparse o repartirse.","Es la cadena más larga del proyecto y está en una sola persona."),
 ("5","Si se recorta alcance.","El margen de cronograma alcanza, pero la sobreasignación de carga no se corrige con margen: es la palanca que queda si las anteriores no bastan."),
],widths=[0.8,6.4,8.8],fs=9.5)

H(doc,"15.4 Nota sobre el método",2)
P(doc,"El análisis se hizo sobre la lista de 257 tareas y el documento de cambios tal como fueron "
  "entregados, sin agregar ni quitar actividades y sin modificar ninguna duración. Se emplearon las mismas "
  "reglas de conversión de tiempos en ambas versiones del análisis, precisamente para que las cifras fueran "
  "comparables entre sí.")
P(doc,"Los cálculos son reproducibles. Cualquier cambio en una duración o en una dependencia se propaga a "
  "las matrices, a las figuras y a las conclusiones repitiendo los recorridos descritos en las secciones 7 "
  "y 8.")

out="/Users/abiudbenitez/Documents/Claude/Projects/vr-bateria-terapeutica/entregables/03_Ruta_Critica.docx"
doc.save(out); print("OK", out)
print("parrafos:",len(doc.paragraphs)," tablas:",len(doc.tables),
      " imagenes:",len(doc.inline_shapes)," secciones:",len(doc.sections))
