# -*- coding: utf-8 -*-
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from neutro import *
from docx import Document
from docx.shared import Cm
import rc257 as R, costos as K
FIG="/Users/abiudbenitez/Documents/Claude/Projects/vr-bateria-terapeutica/entregables/figuras-costos/"
doc = Document('/tmp/_plan_b.docx')

H(doc,"14. Registro de riesgos y respuestas",2)
P(doc,"Cada riesgo lleva su estrategia de respuesta, el responsable de ejecutarla, el disparador que indica "
  "que hay que actuar y el plan de contingencia si el riesgo se materializa. El impacto se expresa en días "
  "hábiles de retraso y en la degradación concreta que produciría sobre el producto.")
if os.path.exists(FIG+"Fig4_Matriz_riesgos.png"):
    doc.add_picture(FIG+"Fig4_Matriz_riesgos.png", width=Cm(14.2)); doc.paragraphs[-1].alignment=C
caption(doc,"Matriz de probabilidad e impacto de los nueve riesgos identificados.")

RESP = {
 "R1": ("Mitigar","Desarrollador VR",
        "El hito de verificación de latencia arroja una medición superior a 30 ms.",
        "Sincronizar el audio contra el reloj del sistema de audio y no contra el ciclo de cuadro. Reducir el tamaño del búfer. Si aun así no se alcanza el umbral, se convoca al director para acotar el alcance del sistema de audio y se documenta la limitación."),
 "R2": ("Mitigar","Responsable de Entorno 3D",
        "Cualquier actividad de la cadena del entorno acumula un día de retraso.",
        "Traslapar la búsqueda de assets con la composición del escenario. Reasignar temporalmente apoyo de otra área. La cadena tiene 17 actividades seriales y concentra la mitad de la ruta crítica, de modo que se vigila semanalmente."),
 "R3": ("Mitigar","Gerente de proyecto",
        "La carga acumulada del área de QA supera el 110 % de su ventana disponible.",
        "Redistribuir documentación, evidencias y manuales entre las ocho áreas, dejando en QA solo pruebas y gestión de errores. Es la medida que más brecha cierra."),
 "R4": ("Mitigar","Responsable de Música y de Sonido",
        "Un recurso incorporado no tiene licencia verificable.",
        "Verificar y registrar la licencia antes de incorporar cualquier recurso. Mantener una lista de recursos sustitutos. Si un recurso debe retirarse, se reemplaza desde esa lista."),
 "R5": ("Transferir","Director de proyecto",
        "No hay confirmación escrita del préstamo antes del inicio de la fase de desarrollo.",
        "Obtener confirmación por escrito de la Facultad. Si no se concreta, se reduce a un solo visor prestado por algún integrante o se solicita apoyo a otra dependencia, con el impacto correspondiente sobre las sesiones simultáneas."),
 "R6": ("Mitigar","Coordinador",
        "Un visor presenta fallas o deja de estar disponible.",
        "Mantener el desarrollo ejecutable también en el simulador de escritorio, de modo que la falta de visor no detenga el trabajo. Reportar la incidencia a la Facultad de inmediato."),
 "R7": ("Mitigar","Diseñador UX-XR",
        "Un participante reporta molestia durante o después de una sesión.",
        "Aplicar los lineamientos de comodidad desde el diseño: sin movimiento impuesto, interfaz a distancia adecuada, sesiones limitadas en duración. Suspender la sesión ante el primer síntoma."),
 "R8": ("Mitigar","Gerente de proyecto",
        "Se detecta un conflicto de integración que requiere rehacer trabajo.",
        "Rama por funcionalidad e integración frecuente. Revisión antes de integrar. Respaldo del repositorio."),
 "R9": ("Aceptar","Director de proyecto",
        "Un integrante comunica que su disponibilidad se reduce.",
        "Es un riesgo inherente a un proyecto académico y no puede eliminarse. Se acepta de forma activa: se reasignan sus actividades entre el resto y, si hace falta, se recorta alcance conforme a la tolerancia declarada."),
}
for rid, desc, prob, dias, alcance, cat in K.RIESGOS:
    est, resp, disp, plan = RESP[rid]
    H(doc, f"{rid} — {desc}", 3)
    table(doc,["Campo","Contenido"],[
     ("Categoría",cat),
     ("Probabilidad",f"{prob:.0%}"),
     ("Impacto en el cronograma",f"{dias:.1f} días hábiles"),
     ("Impacto en el alcance",alcance),
     ("Valor esperado",f"{prob*dias:.2f} días hábiles"),
     ("Estrategia de respuesta",est),
     ("Responsable",resp),
     ("Disparador",disp),
     ("Plan de respuesta y contingencia",plan),
    ],widths=[4.0,12.0],fs=9)

table(doc,["Concepto","Días hábiles"],[
 ("Valor esperado total del registro", f"{K.EMV_DIAS:.2f}"),
 ("Reserva de cronograma disponible", f"{K.RESERVA_CRONO:.2f}"),
 ("Margen", f"{K.RESERVA_CRONO-K.EMV_DIAS:.2f}"),
],widths=[11.0,5.0],fs=10)
P(doc,f"La reserva alcanza, con un margen del {100*(K.RESERVA_CRONO/K.EMV_DIAS-1):.0f} % sobre el valor "
  "esperado. Conviene leerlo con cautela: el valor esperado es un promedio, no un tope. Si se materializaran "
  "simultáneamente R1, R2 y R5, el retraso sumaría quince días y la reserva sería insuficiente.")

doc.save('/tmp/_plan_c.docx'); print("C OK")
