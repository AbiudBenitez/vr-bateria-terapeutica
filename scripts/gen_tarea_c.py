# -*- coding: utf-8 -*-
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from neutro import *
from docx import Document
from docx.shared import Cm
import pmbok as PB
FIG="/Users/abiudbenitez/Documents/Claude/Projects/vr-bateria-terapeutica/tareas/2026-09-14-tabla-pmbok/figuras/"
doc = Document('/tmp/_tarea_b.docx')

H(doc,"3.7 Evolución del estándar: sexta, séptima y octava edición",2)
P(doc,"La estructura descrita hasta aquí corresponde a la sexta edición. Conviene dejar constancia de que "
  "el estándar cambió de forma sustancial después, porque afecta la vigencia de lo que se estudia.")
table(doc,["Aspecto","Sexta edición (2017)","Séptima edición (2021)","Octava edición (2026)"],[
 ("Enfoque","Basado en procesos.","Basado en principios y en la entrega de valor.","Combina principios, dominios y procesos accionables."),
 ("Organización","5 grupos de procesos y 10 áreas de conocimiento.","12 principios y 8 dominios de desempeño.","6 principios y 7 dominios de desempeño."),
 ("Procesos","49, descritos con entradas, herramientas y salidas.","No se describen procesos; el detalle se traslada a un repositorio en línea.","Se reincorporan procesos accionables."),
 ("Extensión","Cerca de 750 páginas.","Cerca de 250 páginas.","Intermedia."),
 ("Uso en este trabajo","Base de la tabla de entradas y salidas, porque es la única edición que las describe.","Se cita como antecedente del cambio de enfoque.","Se cita como estado vigente del estándar."),
],widths=[2.6,4.4,4.4,4.6],fs=9.5)
P(doc,"La séptima edición no invalidó la sexta: el propio instituto mantuvo ambas disponibles, porque los "
  "proyectos predictivos siguen requiriendo el detalle de procesos que solo la sexta ofrece. La octava "
  "edición, publicada el 13 de enero de 2026, parece reconocer esa necesidad al reincorporar procesos.")
P(doc,"Para efectos de esta actividad, que solicita expresamente entradas y salidas por grupo de procesos, "
  "la sexta edición es la referencia obligada.")

doc.add_page_break()
# ================================================================ 4
H(doc,"4. Diagramas",1)
P(doc,"Se presentan tres figuras que apoyan lo expuesto en la sección anterior.")

H(doc,"4.1 Los cinco grupos de procesos y su interacción",2)
if os.path.exists(FIG+"Fig1_Grupos_de_procesos.png"):
    doc.add_picture(FIG+"Fig1_Grupos_de_procesos.png", width=Cm(16.0)); doc.paragraphs[-1].alignment=C
caption(doc,"Fig. 1 — Secuencia de los grupos de procesos, retroalimentación y vigilancia continua.")
P(doc,"La figura muestra tres cosas que las tablas no dejan ver. La primera, que la secuencia no es lineal: "
  "existe un ciclo de retroalimentación de la ejecución hacia la planificación, porque todo cambio aprobado "
  "obliga a replanificar. La segunda, que el monitoreo y control no es una etapa entre la ejecución y el "
  "cierre, sino una actividad que corre en paralelo a todas las demás de principio a fin. La tercera, que "
  "casi la mitad de los procesos, 24 de 49, están en planificación, lo que revela dónde pone el estándar el "
  "peso del esfuerzo.")

H(doc,"4.2 Modelo de entradas, herramientas y salidas",2)
if os.path.exists(FIG+"Fig2_Modelo_ITTO.png"):
    doc.add_picture(FIG+"Fig2_Modelo_ITTO.png", width=Cm(16.0)); doc.paragraphs[-1].alignment=C
caption(doc,"Fig. 2 — Estructura con la que el estándar describe cada uno de sus 49 procesos.")
P(doc,"El encadenamiento que representa la flecha inferior es lo que convierte al conjunto en un sistema y "
  "no en una lista. Si la salida de un proceso está mal construida, el error no se queda ahí: viaja hacia "
  "todos los procesos que la reciben como entrada.")

H(doc,"4.3 Matriz de áreas de conocimiento por grupos de procesos",2)
if os.path.exists(FIG+"Fig3_Matriz_areas_grupos.png"):
    doc.add_picture(FIG+"Fig3_Matriz_areas_grupos.png", width=Cm(16.0)); doc.paragraphs[-1].alignment=C
caption(doc,"Fig. 3 — Ubicación de los 49 procesos en el cruce de las dos dimensiones del método.")
P(doc,"La lectura por columnas confirma la concentración en planificación. La lectura por filas muestra que "
  "solo la gestión de la integración tiene presencia en los cinco grupos, y que la columna de cierre "
  "contiene un único proceso, lo que suele interpretarse erróneamente como que el cierre es poco "
  "importante, cuando lo que indica es que se trata de un acto formal y concentrado.")

doc.add_page_break()
doc.save('/tmp/_tarea_c.docx'); print("C OK")
