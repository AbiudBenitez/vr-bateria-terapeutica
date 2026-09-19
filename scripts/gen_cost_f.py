# -*- coding: utf-8 -*-
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from neutro import *
from docx import Document
import rc257 as R, costos as K, compresion as CP
M = R.V2
def d(x): return f"${x:,.0f}"
def d2(x): return f"${x:,.2f}"
doc = Document('/tmp/_cost_e.docx')
c = K.corte(9.0)

# ================================================================ 14
H(doc,"14. Conclusiones",1)
H(doc,"14.1 Resultado de la estimación",2)
table(doc,["Concepto","Valor"],[
 ("Actividades estimadas","257"),
 ("Esfuerzo total",f"{K.HORAS_TOT:,.0f} horas"),
 ("Tarifa base","1 salario mínimo por hora, "+d2(K.SM_HORA)),
 ("Escala de tarifas","De 1.0 a 2.2 salarios mínimos, según responsabilidad"),
 ("Tarifa media ponderada",d2(K.TARIFA_MED)+f" por hora ({K.TARIFA_MED/K.SM_HORA:.2f} SM)"),
 ("Mano de obra",d(K.MANO_OBRA)),
 ("Costos no laborales","Ninguno. El proyecto no realiza compras"),
 ("Reserva de contingencia monetaria","Ninguna. Sustituida por reserva de cronograma"),
 ("Reserva de cronograma",f"{K.RESERVA_CRONO:.2f} días hábiles"),
 ("Presupuesto hasta la conclusión",d(K.PRESUPUESTO)),
],widths=[9.0,7.0],fs=10)

H(doc,"14.2 Hallazgos",2)
bullets(doc,[
 "El presupuesto se estimó de forma ascendente sobre las 257 actividades, que es la técnica más exacta de "
 "las que reconoce la guía y la única aplicable con el nivel de detalle disponible.",
 "La tarifa se ancla al salario mínimo general vigente, publicado por la CONASAMI, y se gradúa en múltiplos "
 "según la responsabilidad del rol. No se emplearon portales de empleo.",
 "La conversión a tarifa horaria se hace dividiendo el salario mínimo diario entre las ocho horas de la "
 "jornada legal, porque las prácticas profesionales bajo convenio escolar no generan relación laboral ni "
 "días de descanso pagados.",
 "El proyecto no realiza compras: el equipo lo presta la Facultad y los recursos gráficos, sonoros y de "
 "software se resuelven con licencias libres o gratuitas para estudiantes.",
 "Por esa razón ningún riesgo identificado tiene impacto monetario, y la reserva de contingencia se "
 f"sustituye por una reserva de cronograma de {K.RESERVA_CRONO:.2f} días hábiles, suficiente para el valor "
 f"esperado de {K.EMV_DIAS:.2f} días que arroja el análisis de riesgos.",
 f"El área de QA, documentación y gestión concentra el "
 f"{100*K.POR_AREA['QA, documentación y gestión'][2]/K.MANO_OBRA:.0f} % del costo, porque absorbe la "
 "planificación y el seguimiento además de las pruebas y la documentación.",
 f"Al corte de medio curso el índice de desempeño del cronograma es {c['EV']/c['PV']:.3f}, con "
 f"{len(c['terminadas'])} de 257 actividades terminadas.",
])

H(doc,"14.3 Sobre la evolución del presupuesto",2)
P(doc,"Este presupuesto es resultado de tres revisiones sucesivas, y conviene dejar constancia del "
  "razonamiento porque la diferencia entre ellas no es de cálculo sino de supuestos.")
table(doc,["Versión","Supuesto de tarifa","Otros costos","Presupuesto"],[
 ("Primera","Tarifas de mercado tomadas de referencias no oficiales","Equipo, licencias y operación","$559,891"),
 ("Segunda","Observatorio Laboral: profesionistas titulados","Equipo prestado, licencias libres","$294,981"),
 ("Tercera, la vigente","Salario mínimo con escala por responsabilidad, practicantes","Ninguno","$104,013"),
],widths=[2.6,6.0,4.0,3.4],fs=9.5)
P(doc,"La primera corrección obedeció a que los portales de empleo publican el salario ofrecido en vacantes, "
  "que supera sistemáticamente al percibido. La segunda, a reconocer que el equipo no está formado por "
  "profesionistas titulados sino por estudiantes en prácticas, y que el proyecto no compra nada. Ambas "
  "correcciones van en la misma dirección: acercar la estimación a lo que el proyecto realmente es.")

H(doc,"14.4 Recomendaciones",2)
table(doc,["#","Recomendación","Efecto"],[
 ("1","Confirmar por escrito el préstamo de los visores con la Facultad antes de iniciar la fase de desarrollo.",
  "Elimina el riesgo R5, que es el de mayor impacto unitario del registro con seis días hábiles."),
 ("2","Redistribuir las tareas de documentación, evidencias y manuales del área de QA entre las ocho áreas.",
  "Atiende el riesgo R3 y corrige la restricción dominante del proyecto, que es de carga y no de costo."),
 ("3","Asignar de forma explícita los roles de dirección, gerencia y coordinación.",
  "Saca de QA el trabajo de planificación y seguimiento, que hoy está concentrado ahí por omisión."),
 ("4","Comenzar a registrar las horas efectivamente dedicadas a cada actividad.",
  "Sin ese dato no puede calcularse el costo real ni el índice de desempeño del costo, y la mitad del método de valor ganado queda inutilizable."),
 ("5","Verificar y documentar la licencia de cada recurso gráfico y sonoro antes de incorporarlo.",
  "Atiende el riesgo R4, cuya probabilidad es alta al haber sustituido recursos de pago por recursos libres."),
 ("6","Reservar la compresión de la red como respuesta ante la materialización de un riesgo de cronograma.",
  "Convierte una capacidad ya cuantificada en un plan de contingencia concreto."),
],widths=[0.8,8.4,6.8],fs=9.5)

# ================================================================ anexos
doc.add_page_break()
H(doc,"Anexo A. Costeo de las 257 actividades",1)
P(doc,"Cada actividad con su perfil asignado, múltiplo del salario mínimo, tarifa, horas y costo. Es el "
  "nivel más bajo de la estimación ascendente; todas las cifras del documento se obtienen sumando esta "
  "tabla.")
rows=[(k, R.T[k]["desc"], R.T[k]["sig"], K.perfil(k)[:24], f"{K.multiplo(K.perfil(k)):.1f}",
       d2(K.tarifa(k)), f"{K.horas(k):g}", d(K.costo(k))) for k in R.ORD]
rows.append(("","Total","","","","",f"{K.HORAS_TOT:,.0f}",d(K.MANO_OBRA)))
table(doc,["Clave","Actividad","Á.","Perfil","SM","$/h","Horas","Costo"],rows,
      widths=[1.3,5.2,0.7,3.2,0.8,1.3,1.1,1.9],fs=7)

doc.add_page_break()
H(doc,"Anexo B. Pendientes de costo de las actividades críticas",1)
P(doc,"Las actividades de la ruta crítica, ordenadas de la más barata a la más cara de comprimir.")
rows=[(k, R.T[k]["desc"][:52], R.fnum(R.T[k]["dur"]), f"{CP.comprimible(k):.2f}", d(K.costo(k)),
       d(CP.sobrecosto(k)), d(CP.pendiente(k))) for k in CP.CANDIDATAS]
table(doc,["Clave","Actividad","Dur.","Recorte","Costo normal","Sobrecosto","$/día"],rows,
      widths=[1.3,6.2,1.1,1.4,2.2,2.0,1.8],fs=8)

doc.add_page_break()
H(doc,"Anexo C. Fuentes consultadas",1)
table(doc,["Fuente","Institución","Dato empleado"],[
 ("Salarios mínimos generales y profesionales vigentes a partir del 1 de enero de 2026",
  "Comisión Nacional de los Salarios Mínimos","Salario mínimo general de la zona resto del país: "+d2(K.SM_DIARIO)+" diarios. Base de toda la escala de tarifas."),
 ("Programa Jóvenes Construyendo el Futuro","Secretaría del Trabajo y Previsión Social",
  "Apoyo mensual a personas aprendices, equivalente al salario mínimo general. Confirma la referencia para perfiles en formación."),
 ("Ley Federal del Trabajo, artículo 61","Cámara de Diputados del H. Congreso de la Unión",
  "Duración máxima de la jornada diurna: ocho horas. Divisor para obtener la tarifa horaria."),
 ("Ley Federal del Trabajo, artículos 66 a 68","Cámara de Diputados del H. Congreso de la Unión",
  "Jornada extraordinaria. Base del modelo de compresión de la red."),
 ("Observatorio Laboral Mexicano, 2.º trimestre de 2026","Secretaría del Trabajo y Previsión Social, con datos de la ENOE del INEGI",
  "Ingreso de profesionistas titulados. Se emplea solo como contraste en la sección 4.4."),
 ("Guía de los Fundamentos para la Dirección de Proyectos (Guía del PMBOK)","Project Management Institute",
  "Marco de referencia de las áreas de conocimiento 7, 8 y 11."),
],widths=[5.4,4.6,6.0],fs=9)

out="/Users/abiudbenitez/Documents/Claude/Projects/vr-bateria-terapeutica/entregables/Analisis_Costos_Calidad_Riesgos.docx"
doc.save(out); print("OK", out)
print("párrafos:",len(doc.paragraphs)," tablas:",len(doc.tables)," imágenes:",len(doc.inline_shapes))
