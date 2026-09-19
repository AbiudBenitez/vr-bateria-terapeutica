# -*- coding: utf-8 -*-
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from neutro import *
from docx import Document
from docx.shared import Cm
import rc257 as R, costos as K
M = R.V2
def d(x): return f"${x:,.0f}"
FIG="/Users/abiudbenitez/Documents/Claude/Projects/vr-bateria-terapeutica/entregables/figuras-costos/"
doc = Document('/tmp/_cost_c.docx')

# ================================================================ 11
H(doc,"11. Análisis del costo de la calidad",1)
P(doc,"El costo de la calidad reúne todo el esfuerzo dedicado a que el producto cumpla sus requisitos, más "
  "el que se gasta cuando no los cumple. La guía PMBOK lo divide en costos de conformidad, que se invierten "
  "para evitar fallos, y costos de no conformidad, que se pagan cuando el fallo ya ocurrió.")
table(doc,["Categoría","Tipo","Qué comprende en este proyecto"],[
 ("Prevención","Conformidad","Planificación y gestión del proyecto, gestión de riesgos, definición de requisitos y trazabilidad, y registro del avance."),
 ("Evaluación","Conformidad","Planes y ejecución de pruebas, y validación integral con métricas."),
 ("Fallos internos","No conformidad","Gestión y corrección de errores, y correcciones de la batería, la interfaz y la ambientación."),
 ("Fallos externos","No conformidad","No aplica. El prototipo no llega a usuarios finales fuera del entorno de prueba."),
],widths=[2.6,2.4,11.0],fs=9.5)
rows=[]
for cat,gs in K.COQ.items():
    for n,g in enumerate(gs):
        rows.append((cat if n==0 else "", g, K.GRUPO[g][:28], d(K.costo_grupo(g))))
    rows.append(("", f"Subtotal {cat.lower()}", "", d(K.COQ_TOT[cat])))
rows.append(("Total del costo de la calidad","","",d(sum(K.COQ_TOT.values()))))
table(doc,["Categoría","Grupo","Perfil","Costo"],rows,widths=[3.4,1.8,6.4,2.8],fs=9)
table(doc,["Categoría","Costo","% de la mano de obra","% del costo de la calidad"],
 [(c, d(v), f"{100*v/K.MANO_OBRA:.1f} %", f"{100*v/sum(K.COQ_TOT.values()):.1f} %") for c,v in K.COQ_TOT.items()]+
 [("Total", d(sum(K.COQ_TOT.values())), f"{100*sum(K.COQ_TOT.values())/K.MANO_OBRA:.1f} %","100.0 %")],
 widths=[4.4,3.2,4.2,4.2],fs=10)
P(doc,f"El costo de la calidad representa el {100*sum(K.COQ_TOT.values())/K.MANO_OBRA:.1f} % de la mano de "
  "obra. Es una proporción alta comparada con la referencia habitual de la industria del software, que ronda "
  "el 15 %, y tiene una explicación concreta: en un proyecto académico la documentación es en sí misma un "
  "entregable evaluable, de modo que el esfuerzo de prevención incluye tareas que en un proyecto comercial "
  "no existirían.")
P(doc,"La proporción entre categorías sí es sana. Se invierte más en prevención que en corregir fallos, "
  f"{d(K.COQ_TOT['Prevención'])} contra {d(K.COQ_TOT['Fallos internos'])}, que es la relación que se busca: "
  "cuesta menos evitar un defecto que repararlo.")

doc.add_page_break()
# ================================================================ 12
H(doc,"12. Análisis de riesgos",1)
P(doc,"Esta sección identifica los riesgos del proyecto, los evalúa de forma cualitativa y cuantitativa, y "
  "obtiene de ahí la reserva de cronograma. El tratamiento de cada riesgo, con su estrategia de respuesta, "
  "responsable y plan de contingencia, se desarrolla en el Plan de Gestión de los Riesgos.")

H(doc,"12.1 Por qué el impacto se mide en días y no en pesos",2)
P(doc,"En la mayoría de los proyectos el impacto de un riesgo se expresa en dinero, porque materializarse "
  "significa comprar algo que no estaba previsto, pagar horas adicionales o afrontar una penalización. Nada "
  "de eso ocurre aquí: el proyecto no compra, no contrata y no tiene cliente que penalice.")
P(doc,"Lo que sí puede perder es tiempo y alcance. Por eso el impacto de cada riesgo se valora en días "
  "hábiles de retraso y en la degradación concreta que produciría sobre el producto, y la reserva que se "
  "constituye es de cronograma, no de dinero. Es la aplicación del mismo método sobre la variable que en "
  "este proyecto sí está en riesgo.")

H(doc,"12.2 Escalas de valoración",2)
table(doc,["Nivel","Probabilidad","Impacto en cronograma","Impacto en alcance"],[
 ("Muy bajo","10 %","Alrededor de 1 día hábil","Ningún entregable afectado"),
 ("Bajo","30 %","Alrededor de 2 días","Un entregable secundario se degrada"),
 ("Medio","50 %","Alrededor de 3 días","Un entregable principal se degrada"),
 ("Alto","70 %","Alrededor de 5 días","Se pierde un entregable secundario"),
 ("Muy alto","90 %","8 días o más","Se pierde un entregable principal"),
],widths=[2.0,2.2,4.4,7.4],fs=9.5)

H(doc,"12.3 Registro de riesgos y análisis cuantitativo",2)
P(doc,"El valor esperado de cada riesgo es el producto de su probabilidad por su impacto en días. La suma "
  "determina la reserva de cronograma necesaria.")
rows=[(r[0], r[1], r[5], f"{r[2]:.0%}", f"{r[3]:.1f} d", f"{r[2]*r[3]:.2f} d", r[4]) for r in K.RIESGOS]
rows.append(("","Valor esperado total","","","",f"{K.EMV_DIAS:.2f} d",""))
table(doc,["Id","Riesgo","Categoría","Prob.","Impacto","Valor esp.","Impacto en alcance"],rows,
      widths=[0.8,4.8,1.8,1.1,1.3,1.4,4.8],fs=8.5)

H(doc,"12.4 Matriz de probabilidad e impacto",2)
if os.path.exists(FIG+"Fig4_Matriz_riesgos.png"):
    doc.add_picture(FIG+"Fig4_Matriz_riesgos.png", width=Cm(14.6)); doc.paragraphs[-1].alignment=C
caption(doc,"Fig. 4 — Matriz de probabilidad e impacto de los nueve riesgos identificados.")

H(doc,"12.5 Suficiencia de la reserva",2)
table(doc,["Concepto","Días hábiles"],[
 ("Valor esperado del impacto de los nueve riesgos", f"{K.EMV_DIAS:.2f}"),
 ("Reserva de cronograma disponible", f"{K.RESERVA_CRONO:.2f}"),
 ("Margen", f"{K.RESERVA_CRONO-K.EMV_DIAS:.2f}"),
],widths=[11.0,5.0],fs=10)
P(doc,f"La reserva alcanza, con un margen del {100*(K.RESERVA_CRONO/K.EMV_DIAS-1):.0f} % sobre el valor "
  "esperado. Conviene leerlo con cautela: el valor esperado es un promedio, no un tope. Si se materializaran "
  "simultáneamente los tres riesgos de mayor impacto —el incumplimiento del umbral de latencia, el retraso "
  "del entorno tridimensional y la falta del préstamo de los visores— el retraso sumaría 15 días y la "
  "reserva sería insuficiente.")
P(doc,"Los tres riesgos de mayor valor esperado son la sobrecarga del área de QA, el retraso de la cadena "
  "del entorno tridimensional y el incumplimiento del umbral de latencia. Los dos primeros son de recursos y "
  "cronograma, y se atienden redistribuyendo trabajo; el tercero es técnico y se atiende con verificación "
  "temprana.")
P(doc,"El riesgo R5, que la Facultad no concrete el préstamo de los visores, merece atención particular: es "
  "el de mayor impacto unitario, seis días hábiles, y además reduciría las pruebas con usuarios a una sola "
  "sesión. Es la contrapartida de haber excluido el equipo del presupuesto.")

doc.add_page_break()
doc.save('/tmp/_cost_d.docx'); print("D OK")
